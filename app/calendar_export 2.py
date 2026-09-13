from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


NAVY = "315C9B"
PEACH = "F1D0BC"
INK = "18201D"
MUTED = "6C756F"
LINE = "DDE2DC"
WHITE = "FFFFFF"


def build_calendar_event_workbook(payload: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bakım ve Mesai"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A6"
    sheet.merge_cells("A1:J2")
    title = sheet["A1"]
    title.value = "SELSA  ·  BAKIM, MESAİ VE VARDİYA GEÇMİŞİ"
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.font = Font(name="Aptos Display", size=17, color=WHITE, bold=True)
    title.alignment = Alignment(vertical="center")
    sheet.merge_cells("A3:J3")
    sheet["A3"] = f"Filtrelenmiş kayıtlar  ·  {len(payload.get('rows', []))} kayıt  ·  Oluşturulma: {payload.get('generatedAt', '')}"
    sheet["A3"].font = Font(name="Aptos", size=9, color=MUTED)

    headers = ["Tür", "Proses", "İş merkezi", "İş merkezi adı", "Açıklama", "Başlangıç", "Bitiş", "Vardiya", "Durum", "Tamamlanma kaydı"]
    widths = [20, 15, 15, 34, 38, 14, 14, 11, 14, 24]
    thin = Side(style="thin", color=LINE)
    for column, (header, width) in enumerate(zip(headers, widths), 1):
        cell = sheet.cell(5, column, header)
        cell.fill = PatternFill("solid", fgColor=PEACH)
        cell.font = Font(name="Aptos", size=9, color="56301E", bold=True)
        cell.alignment = Alignment(vertical="center")
        cell.border = Border(bottom=thin)
        sheet.column_dimensions[cell.column_letter].width = width

    labels = {"maintenance": "Planlı bakım", "overtime": "Mesai", "shift-change": "Vardiya değişikliği"}
    process_labels = {"turning": "Torna", "drilling": "Delme", "deburring": "Çapak Alma", "gkm": "GKM"}
    for row_index, row in enumerate(payload.get("rows", []), 6):
        values = [
            labels.get(row.get("eventType"), row.get("eventType", "")),
            process_labels.get(row.get("process"), row.get("process", "")),
            row.get("workCenterId", ""),
            row.get("workCenterName", ""),
            row.get("name", ""),
            row.get("startDate", ""),
            row.get("endDate", ""),
            row.get("shiftCount", 0),
            "Tamamlandı" if row.get("status") == "completed" else "Aktif",
            row.get("completedAt", ""),
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column, value)
            cell.fill = PatternFill("solid", fgColor=WHITE)
            cell.font = Font(name="Aptos", size=9, color=INK, bold=column in {1, 3, 9})
            cell.alignment = Alignment(vertical="top", wrap_text=column in {4, 5})
            cell.border = Border(bottom=thin)
        sheet.row_dimensions[row_index].height = 24

    if not payload.get("rows"):
        sheet.merge_cells("A6:J7")
        sheet["A6"] = "Seçili filtrelerle eşleşen kayıt bulunamadı."
        sheet["A6"].font = Font(name="Aptos", size=10, color=MUTED, italic=True)
        sheet["A6"].alignment = Alignment(horizontal="center", vertical="center")
    else:
        sheet.auto_filter.ref = f"A5:J{sheet.max_row}"

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
