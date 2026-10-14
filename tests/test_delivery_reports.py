import hashlib
import json
import subprocess
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import delivery_reports as reports
from app import main


@pytest.fixture
def saved(monkeypatch):
    today = datetime.now(ZoneInfo("Europe/Istanbul")).date().isoformat()
    seed = {"orders": [], "products": [], "preferences": [], "machines": [], "manualBatches": [],
            "productionHistory": [{"id": "shipment", "product": "PART-A", "inventoryStatus": "delivered", "completedQuantity": 80, "deliveredAt": today}]}
    state = {"seed": seed, "updatedAt": "2026-09-16T10:00:00Z"}
    monkeypatch.setattr(reports, "planning_state", lambda: state)
    monkeypatch.setattr(main, "planning_state", lambda: state)
    return state, {"startDate": today, "endDate": today}


@pytest.fixture
def client():
    with TestClient(main.app) as client:
        token = client.post("/api/auth/login", json={"password": "test-password"}).json()["token"]
        client.headers["Authorization"] = f"Bearer {token}"
        yield client


def test_shared_runtime_preview_and_workbook_are_from_saved_state(saved, client):
    state, scope = saved
    request = {"expectedRevision": state["updatedAt"], **scope}
    response = client.post("/api/delivery-plan/preview", json=request)
    assert response.status_code == 200, response.text
    snapshot = response.json()
    assert snapshot["report"]["total"] == 80
    assert snapshot["report"]["remaining"] == 0
    assert snapshot["report"]["complete"] is True
    response = client.post("/api/delivery-plan/export", json={**request, "digest": snapshot["digest"], "calculationDate": snapshot["calculationDate"]})
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["Teslimat Planı"]
    assert sheet["B7"].value == "PART-A"
    assert sum(sheet.cell(7, col).value for col in range(12, sheet.max_column + 1) if type(sheet.cell(7, col).value) is int) == 80
    assert workbook["Rapor doğrulaması"]["B1"].value == state["updatedAt"]
    assert workbook["Rapor doğrulaması"]["B5"].value == snapshot["digest"]


def test_old_client_totals_and_extra_input_cannot_bypass_authority(saved, client):
    state, scope = saved
    assert client.post("/api/delivery-plan/export", json={**scope, "weeks": [], "rows": []}).status_code == 422
    assert client.post("/api/delivery-plan/preview", json={"expectedRevision": state["updatedAt"], **scope, "seed": {}}).status_code == 422
    assert client.post("/api/delivery-plan/preview", json={"expectedRevision": state["updatedAt"], **scope, "startDate": "2026-02-30"}).status_code == 422


def test_revision_and_digest_conflicts_require_new_preview(saved, client):
    state, scope = saved
    request = {"expectedRevision": state["updatedAt"], **scope}
    snapshot = client.post("/api/delivery-plan/preview", json=request).json()
    export = {**request, "digest": snapshot["digest"], "calculationDate": snapshot["calculationDate"]}
    assert client.post("/api/delivery-plan/export", json={**export, "digest": "0" * 64}).status_code == 409
    assert client.post("/api/delivery-plan/export", json={**export, "calculationDate": "2000-01-01"}).status_code == 409
    state["updatedAt"] = "2026-09-16T11:00:00Z"
    assert client.post("/api/delivery-plan/export", json=export).status_code == 409


def test_change_during_runtime_is_rejected(saved, monkeypatch):
    state, scope = saved
    calculate = reports.run_delivery_runtime
    def racing(*args):
        result = calculate(*args)
        state["updatedAt"] = "new"
        return result
    monkeypatch.setattr(reports, "run_delivery_runtime", racing)
    with pytest.raises(reports.DeliveryReportError, match="sırasında"):
        reports.report_snapshot(state["updatedAt"], scope)


def test_explicit_hold_is_separate_and_does_not_block_export(saved, client):
    """HOLD-DEL01: intentional holds are excluded; other shipments remain downloadable."""
    import copy
    state, scope = saved
    state["seed"]["wipLots"] = [{"id": "held", "workOrder": "320-2600726", "product": "PART-B", "availableQuantity": 1920, "stage": "on-hold", "completedSteps": []}]
    before = copy.deepcopy(state)
    request = {"expectedRevision": state["updatedAt"], **scope}
    response = client.post("/api/delivery-plan/preview", json=request)
    assert response.status_code == 200, response.text
    snapshot = response.json()
    assert snapshot["report"]["heldQuantity"] == 1920
    assert snapshot["report"]["held"][0]["workOrder"] == "320-2600726"
    assert snapshot["report"]["undatedQuantity"] == 0
    assert snapshot["report"]["total"] == 80
    assert snapshot["report"]["remaining"] == 0
    assert snapshot["report"]["complete"] is True
    response = client.post("/api/delivery-plan/export", json={**request, "digest": snapshot["digest"], "calculationDate": snapshot["calculationDate"]})
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    held = workbook["Blokeli şarjlar"]
    assert held["A6"].value == "320-2600726"
    assert held["B6"].value == "PART-B"
    assert held["C6"].value == "held"
    assert held["D6"].value == held["D3"].value == 1920
    assert workbook["Rapor doğrulaması"]["B7"].value == 1920
    assert workbook["Rapor doğrulaması"]["B5"].value == snapshot["digest"]
    assert state == before


def test_unclassified_material_still_blocks_certification_even_with_a_hold(saved, client):
    state, scope = saved
    state["seed"]["wipLots"] = [
        {"id": "held", "product": "PART-B", "availableQuantity": 1920, "stage": "on-hold", "completedSteps": []},
        {"id": "unknown", "product": "PART-C", "availableQuantity": 120, "stage": "unclassified", "completedSteps": []},
    ]
    request = {"expectedRevision": state["updatedAt"], **scope}
    response = client.post("/api/delivery-plan/preview", json=request)
    assert response.status_code == 200, response.text
    snapshot = response.json()
    assert snapshot["report"]["heldQuantity"] == 1920
    assert snapshot["report"]["undatedQuantity"] == 120
    assert snapshot["report"]["total"] == 80
    assert snapshot["report"]["complete"] is False
    assert client.post("/api/delivery-plan/export", json={**request, "digest": snapshot["digest"], "calculationDate": snapshot["calculationDate"]}).status_code == 422


def test_runtime_timeout_is_bounded_and_does_not_fallback(saved, monkeypatch):
    state, scope = saved
    def timeout(*args, **kwargs):
        assert kwargs["timeout"] == 20
        assert "NODE_OPTIONS" not in kwargs["env"]
        assert "shell" not in kwargs
        raise subprocess.TimeoutExpired(args[0], kwargs["timeout"])
    monkeypatch.setattr(reports.subprocess, "run", timeout)
    with pytest.raises(reports.DeliveryReportError) as error:
        reports.report_snapshot(state["updatedAt"], scope)
    assert error.value.status == 503


def test_runtime_output_schema_and_bundle_integrity_fail_closed(saved, monkeypatch, tmp_path):
    state, scope = saved
    def malformed(*args, **kwargs):
        kwargs["stdout"].write(b'{"error":"", "complete":true}')
    monkeypatch.setattr(reports.subprocess, "run", malformed)
    with pytest.raises(reports.DeliveryReportError):
        reports.report_snapshot(state["updatedAt"], scope)
    runtime = tmp_path / "delivery-runtime.cjs"
    runtime.write_text("invalid")
    runtime.with_suffix(".cjs.json").write_text(json.dumps({"runtimeSha256": hashlib.sha256(b"different").hexdigest()}))
    monkeypatch.setattr(reports, "RUNTIME", runtime)
    with pytest.raises(reports.DeliveryReportError) as error:
        reports.report_snapshot(state["updatedAt"], scope)
    assert error.value.status == 503


def test_export_scope_change_requires_fresh_preview_and_download_preserves_state(saved, client):
    """DEL-EXP01: changed dates cannot reuse a verified file; export is read-only."""
    import copy
    from datetime import date, timedelta
    state, scope = saved
    before = copy.deepcopy(state)
    request = {"expectedRevision": state["updatedAt"], **scope}
    snapshot = client.post("/api/delivery-plan/preview", json=request).json()
    changed = {**request, "endDate": (date.fromisoformat(scope["endDate"]) + timedelta(days=7)).isoformat()}
    stale = client.post("/api/delivery-plan/export", json={**changed, "digest": snapshot["digest"], "calculationDate": snapshot["calculationDate"]})
    assert stale.status_code == 409
    fresh = client.post("/api/delivery-plan/preview", json=changed).json()
    exported = client.post("/api/delivery-plan/export", json={**changed, "digest": fresh["digest"], "calculationDate": fresh["calculationDate"]})
    assert exported.status_code == 200, exported.text
    assert exported.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert f'Teslimat_Plani_{changed["startDate"]}_{changed["endDate"]}.xlsx' in exported.headers["content-disposition"]
    assert exported.content.startswith(b"PK")
    workbook = load_workbook(BytesIO(exported.content))
    assert workbook["Rapor doğrulaması"]["B5"].value == fresh["digest"]
    assert workbook.calculation.fullCalcOnLoad is True
    assert state == before


def test_cnc_plus_three_export_has_separate_authority_and_preserves_actual_shipments(saved, client):
    from datetime import date, timedelta
    state, scope = saved
    today = date.fromisoformat(scope["startDate"])
    serial = (today - date(1899, 12, 30)).days
    state["seed"]["manualBatches"] = [{"id": "cnc-forecast", "sequence": 1, "product": "PART-A", "machineId": "C-01", "quantity": 100,
        "orderQuantity": 100, "start": serial - 4, "end": serial - 3 + 0.5, "status": "planned", "diameter": "20", "batchSize": 100,
        "dailyRate": 100, "availableStart": serial - 4, "planEnd": serial + 7, "setupHours": 0, "runHours": 12}]
    state["seed"]["wipLots"] = [{"id": "cnc-held", "workOrder": "HOLD-CNC", "product": "PART-B", "availableQuantity": 1920, "stage": "on-hold", "completedSteps": []}]
    scope = {"startDate": scope["startDate"], "endDate": (today + timedelta(days=7)).isoformat()}
    request = {"expectedRevision": state["updatedAt"], **scope}
    before = json.dumps(state, sort_keys=True)
    preview = client.post("/api/delivery-plan/cnc-plus-three/preview", json=request)
    assert preview.status_code == 200, preview.text
    snapshot = preview.json()
    assert snapshot["report"]["heldQuantity"] == 1920
    assert snapshot["report"]["complete"] is True
    assert snapshot["report"]["remaining"] == 100
    assert snapshot["report"]["total"] == 180
    assert snapshot["report"]["events"][0]["readyAt"] == serial + 0.5
    response = client.post("/api/delivery-plan/cnc-plus-three/export", json={**request, "digest": snapshot["digest"], "calculationDate": snapshot["calculationDate"]})
    assert response.status_code == 200, response.text
    assert "CNC_3_Gun" in response.headers["Content-Disposition"]
    workbook = load_workbook(BytesIO(response.content))
    assert workbook["Blokeli şarjlar"]["D6"].value == 1920
    assert "CNC bitişi + 3 takvim günü" in workbook["Rapor doğrulaması"]["B6"].value
    assert json.dumps(state, sort_keys=True) == before
    ordinary = client.post("/api/delivery-plan/preview", json=request).json()
    assert ordinary["digest"] != snapshot["digest"]
    assert client.post("/api/delivery-plan/cnc-plus-three/export", json={**request, "digest": ordinary["digest"], "calculationDate": ordinary["calculationDate"]}).status_code == 409
    assert client.post("/api/delivery-plan/export", json={**request, "digest": snapshot["digest"], "calculationDate": snapshot["calculationDate"]}).status_code == 409
    assert client.post("/api/delivery-plan/cnc-plus-three/preview", json={**request, "expectedRevision": "stale"}).status_code == 409
    state["seed"]["manualBatches"][0]["end"] = 0
    invalid = client.post("/api/delivery-plan/cnc-plus-three/preview", json=request)
    assert invalid.status_code == 200
    assert invalid.json()["report"]["complete"] is False
    bad = invalid.json()
    assert client.post("/api/delivery-plan/cnc-plus-three/export", json={**request, "digest": bad["digest"], "calculationDate": bad["calculationDate"]}).status_code == 422


@pytest.mark.parametrize("corruption", ["quantity", "missing", "unheld"])
def test_runtime_rejects_invalid_held_exclusions(saved, monkeypatch, corruption):
    state, scope = saved
    state["seed"]["wipLots"] = [{"id": "held", "product": "PART-B", "availableQuantity": 1920, "stage": "on-hold", "completedSteps": []}]
    report, _ = reports.run_delivery_runtime(state["seed"], scope, scope["startDate"])
    if corruption == "quantity":
        report["heldQuantity"] = 1919
    elif corruption == "missing":
        report["held"] = []
        report["heldQuantity"] = 0
    else:
        report["held"][0]["id"] = "not-held"
    def invalid(*args, **kwargs):
        kwargs["stdout"].write(json.dumps(report).encode())
    monkeypatch.setattr(reports.subprocess, "run", invalid)
    with pytest.raises(reports.DeliveryReportError) as error:
        reports.run_delivery_runtime(state["seed"], scope, scope["startDate"])
    assert error.value.status == 503
