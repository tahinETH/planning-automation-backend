import copy
from concurrent.futures import ThreadPoolExecutor

import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.config import settings
from app import database
from app.auth import create_test_session
from app.production_area import production_area
from app.ravi.tools import Inspect, inspect_state


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setitem(settings.__dict__, "database_path", tmp_path / "staging.sqlite")
    monkeypatch.setitem(settings.__dict__, "upload_dir", tmp_path / "uploads")
    with TestClient(app) as client:
        yield client


def headers(area="torna", role="admin"):
    return {"Authorization": f"Bearer {create_test_session(settings.admin_password, role)}", "X-Production-Area": area}


def seed(area, quantity):
    return {"productionArea": area, "orders": [{"id": "same-order", "product": "same-product", "quantity": quantity, "dueDate": "2026-10-01"}], "products": [], "machines": [], "preferences": []}


def save(client, area, quantity, version=""):
    return client.put("/api/planning-state", headers=headers(area), json={"seed": seed(area, quantity), "expectedUpdatedAt": version})


def test_cf_p01_separate_plans_orders_versions_history_and_scenarios(client):
    first = save(client, "torna", 480)
    assert first.status_code == 200, first.text
    original = client.get("/api/planning-state", headers=headers()).json()
    created = save(client, "cubuk-filtre", 16000)
    assert created.status_code == 200, created.text
    assert created.headers["x-production-area"] == "cubuk-filtre"
    assert client.get("/api/planning-state", headers=headers()).json() == original
    assert client.get("/api/planning-state", headers=headers("cubuk-filtre")).json()["seed"]["orders"][0]["quantity"] == 16000
    for area in ["torna", "cubuk-filtre"]:
        payload = {"id": "same-scenario", "name": area, "createdAt": "2026-09-22T12:00:00Z", "notes": "", "seed": seed(area, 480 if area == "torna" else 16000), "result": {"batches": [], "summary": {}}}
        assert client.post("/api/scenarios", headers=headers(area), json=payload).status_code == 201
        assert client.get("/api/scenarios/same-scenario", headers=headers(area)).json()["name"] == area
        assert len(client.get("/api/planning-state/history", headers=headers(area)).json()) == 1
    deleted = client.delete("/api/scenarios/same-scenario", headers=headers("cubuk-filtre"))
    assert deleted.status_code == 200
    assert client.get("/api/scenarios/same-scenario", headers=headers()).status_code == 200
    assert save(client, "cubuk-filtre", 16001, original["updatedAt"]).status_code == 409
    assert client.get("/api/planning-state", headers=headers()).json() == original


def test_cf_p02_wrong_area_and_legacy_payload_rejected_before_write(client):
    for selected, incoming in [("cubuk-filtre", "torna"), ("torna", "cubuk-filtre")]:
        response = client.put("/api/planning-state", headers=headers(selected), json={"seed": seed(incoming, 10), "expectedUpdatedAt": ""})
        assert response.status_code == 422, response.text
        payload = {"id": "bad", "name": "bad", "createdAt": "2026-09-22", "seed": seed(incoming, 10), "result": {}}
        assert client.post("/api/scenarios", headers=headers(selected), json=payload).status_code == 400
        assert client.get("/api/planning-state", headers=headers(selected)).json() is None
    assert client.get("/api/planning-state", headers=headers("../production")).status_code == 400
    assert client.post("/api/production-sync/pull", headers=headers("cubuk-filtre")).status_code == 409
    assert client.get("/api/planning-state", headers={"X-Production-Area": "cubuk-filtre"}).status_code == 401


def test_cf_p03_request_scope_is_concurrent_and_roles_remain_server_owned(client):
    assert save(client, "torna", 480).status_code == 200
    assert save(client, "cubuk-filtre", 16000).status_code == 200
    def read(area):
        return client.get("/api/planning-state", headers=headers(area)).json()["seed"]["orders"][0]["quantity"]
    areas = ["torna", "cubuk-filtre"] * 10
    with ThreadPoolExecutor(max_workers=4) as executor:
        assert list(executor.map(read, areas)) == [480, 16000] * 10
    assert production_area.get() == "torna"
    assert client.get("/api/me", headers=headers("cubuk-filtre", "user")).json()["role"] == "user"
    response = client.put("/api/planning-state", headers=headers("cubuk-filtre", "user"), json={"seed": {**seed("cubuk-filtre", 16000), "holidays": [{"serial": 47000}]}, "expectedUpdatedAt": client.get("/api/planning-state", headers=headers("cubuk-filtre")).json()["updatedAt"]})
    assert response.status_code == 403


def test_cf_p04_production_environment_does_not_expose_filter_area(client, monkeypatch):
    monkeypatch.setitem(settings.__dict__, "app_env", "production")
    assert client.get("/api/planning-state", headers=headers("cubuk-filtre")).status_code == 400
    assert client.get("/api/planning-state", headers=headers()).status_code == 200


def test_cf_r01_ravi_reads_selected_area_and_offers_only_its_navigation(client):
    from app.ravi import knowledge
    assert save(client, "torna", 480).status_code == 200
    assert save(client, "cubuk-filtre", 16000).status_code == 200
    for area, expected in [("torna", 480), ("cubuk-filtre", 16000)]:
        token = production_area.set(area)
        try:
            result = inspect_state(database.planning_state(), Inspect(collection="summary"))
            assert result["productionArea"] == area
            assert result["orderQuantity"] == expected
            targets = [t["id"] for t in knowledge.navigation_targets(False)]
            assert ("filter-production-data" in targets) == (area == "cubuk-filtre")
            assert "settings" not in targets
        finally:
            production_area.reset(token)


@pytest.mark.parametrize("process,label", [("gtm","GTM"),("diameter-grinding","Çap Taşlama"),("form-grinding","Form Taşlama"),("measuring","Ölçme"),("milling","Freze"),("final-inspection","Final Kontrol"),("filter-visual","Göz Kontrol")])
def test_cf_e01_each_new_process_can_download_an_actual_excel_response(client, process, label):
    from io import BytesIO
    from openpyxl import load_workbook
    payload = {"generatedAt":"2026-09-22T12:00:00Z","planState":"empty","dirty":False,"selectedProcess":process,
        "summary":dict.fromkeys(["demand","planned","plannedBatchCount","unplannedBatchCount","activeMachineCount","machineCount","machinesWithWorkCount","plannedMachineCount"],0),
        "machines":[],"operationPlans":[{"process":process,"label":label,"totalQuantity":0,"totalJobCount":0,"resources":[]}],"findings":[]}
    response = client.post('/api/general-overview/export',headers=headers('cubuk-filtre'),json=payload)
    assert response.status_code == 200, response.text
    assert response.headers['content-disposition'].isascii()
    assert load_workbook(BytesIO(response.content)).sheetnames


def test_cf_e02_material_export_uses_product_scrap_once_and_retains_legacy_fallback():
    from decimal import Decimal
    from app.overview_export import _material_amounts
    row={"quantity":16000,"unitWeightGrams":3.33,"scrapPercent":2.4}
    assert _material_amounts(row, 99)[2] == Decimal('54.55872')
    assert _material_amounts({**row,"scrapPercent":None}, 2)[2] == Decimal('54.3456')
