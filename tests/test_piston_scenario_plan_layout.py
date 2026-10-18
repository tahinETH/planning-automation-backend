"""Piston readable views preserve typed business values and backup import."""
from datetime import datetime
from io import BytesIO
from openpyxl import load_workbook
from app.data_package import build_data_package, parse_data_package


def test_piston_plan_layout_keeps_dates_quantities_and_recoverable_payload():
    records = [{"id": "scenario-layout", "name": "Aynı ad", "notes": "Uzun açıklama " * 25, "result": {"summary": {"demand": 600, "planned": 600}}}]
    message = "Eksik rota bilgisini ürün ve şarj kaydıyla kontrol edin. " * 12
    plans = {"engine": "a" * 64, "computedAt": "2026-10-18T12:00:00Z", "scenarios": [{"id": "scenario-layout", "name": "Aynı ad", "status": "ok", "rows": {"turning": [{"id": "charge", "quantity": 600, "start": 46313.5, "end": 46314.5, "setupHours": 1.25, "status": "Planlandı"}]}, "warnings": [{"code": "missing-route", "message": message}]}]}
    content = build_data_package("scenarios", records, plans)
    book = load_workbook(BytesIO(content))
    assert parse_data_package(content, "scenarios") == records
    assert book["__SelsaVeri"].sheet_state == "veryHidden"
    for name in ["Torna Planı", "Delme Planı", "Çapak Alma Planı", "GKM Planı"]:
        sheet = book[name]
        assert sheet.freeze_panes == "C2"
        assert sheet.auto_filter.ref == sheet.dimensions
        assert sheet.cell(1, 12).value == "Setup (saat)"
        assert sheet["P2"].alignment.wrap_text
    turning = book["Torna Planı"]
    assert isinstance(turning["I2"].value, datetime)
    assert turning["I2"].number_format == "dd.mm.yyyy hh:mm"
    assert turning["G2"].value == 600
    assert turning["L2"].value == 1.25
    assert book["Senaryolar"]["H2"].value == "scenario-layout"
    warning = book["Plan Uyarıları"]
    assert warning["F2"].value == message
    assert warning["F2"].alignment.wrap_text
    assert warning.row_dimensions[2].height >= 80
    assert book["Bilgi"]["B7"].alignment.wrap_text


def test_filter_scenario_summary_layout_and_payload_stay_unchanged():
    records = [{"id": "filter-id", "name": "CF", "seed": {"productionArea": "cubuk-filtre"}, "result": {}}]
    content = build_data_package("scenarios", records)
    book = load_workbook(BytesIO(content))
    assert book["Senaryolar"].max_column == 7
    assert book["Senaryolar"].freeze_panes == "A2"
    assert "Torna Planı" not in book.sheetnames
    assert parse_data_package(content, "scenarios") == records
