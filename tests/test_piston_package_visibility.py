"""T18: a parameter backup must also be readable without the hidden payload."""

from io import BytesIO

from openpyxl import load_workbook
from fastapi.testclient import TestClient

from app.data_package import build_data_package, parse_data_package
from app.auth import create_test_session
from app.main import app


def piston_parameters():
    return {
        "productionArea": "torna",
        "products": [{"product": "PISTON-T18", "setupFamily": "piston", "batchSize": 1920}],
        "preferences": [{"key": "PISTON-T18", "machines": ["C-01"]}],
        "restrictions": [],
        "machineRates": [{"id": "C-01", "rates": {"PISTON-T18": 640}}],
        "processMasterData": {
            "version": 1,
            "source": {"kind": "one-time-foundation", "files": [], "calculatedOnce": ""},
            "machineCodeStandard": [],
            "resources": [
                {"id": resource, "name": label, "process": process, "active": True, "defaultShifts": 2}
                for process, resource, label in [
                    ("drilling", "D-01", "Delme"),
                    ("deburring", "CA-01", "Çapak Alma"),
                    ("gkm", "GKM-01", "GKM"),
                ]
            ],
            "products": [{
                "product": "PISTON-T18", "family": "piston",
                "route": ["turning", "drilling", "deburring", "gkm"],
                "processes": [
                    {"process": process, "sequence": sequence, "waitWorkdaysBefore": 0,
                     "setupKey": "T18", "resourcePriority": [resource],
                     "unitsPerShift": {resource: rate}, "setupMinutesByResource": {resource: 17}}
                    for sequence, (process, resource, rate) in enumerate([
                        ("drilling", "D-01", 803),
                        ("deburring", "CA-01", 607),
                        ("gkm", "GKM-01", 409),
                    ], 1)
                ],
            }],
        },
    }


def test_t18_downstream_parameters_survive_backup_import():
    data = piston_parameters()
    assert parse_data_package(build_data_package("parameters", data), "parameters") == data


def test_t18_downstream_product_resource_rates_are_readable_in_visible_rows():
    workbook = load_workbook(BytesIO(build_data_package("parameters", piston_parameters())), read_only=True)
    rows = [
        tuple(row)
        for sheet in workbook if sheet.sheet_state == "visible"
        for row in sheet.iter_rows(min_row=2, values_only=True)
    ]
    # Do not fix sheet names. Each visible assignment must identify its product,
    # resource and shift quantity together; hidden JSON is not a readable table.
    for resource, rate in [("D-01", 803), ("CA-01", 607), ("GKM-01", 409)]:
        assert any("PISTON-T18" in row and resource in row and rate in row for row in rows), (
            f"Visible parameter rows must show PISTON-T18, {resource}, and {rate} pieces/shift"
        )


def test_t18_one_workbook_contains_parameters_settings_and_readable_flow_values():
    settings = {
        "setupSettings": {"shiftHours": 8, "sameDiameterHours": 1},
        "machines": [{"id": "C-01", "name": "T18 machine", "active": True, "shiftFactor": 2}],
        "holidays": [], "calendarEvents": [],
        "processWaitSettings": {"drilling": {"enabled": True, "workdays": 2}},
        "flowSettings": {"enabled": True, "defaultShifts": 2, "rules": [
            {"id": "t18-wash", "description": "T18 washing", "family": "piston", "from": "deburring", "to": "gkm", "mode": "static", "minutes": 480, "active": True},
            {"id": "t18-gap", "description": "T18 finish gap", "family": "piston", "from": "turning", "to": "drilling", "mode": "dynamic", "minutes": 120, "active": True},
        ]},
        "productWeights": [{"product": "PISTON-T18", "productName": "T18 piston", "materialCode": "MAT-T18", "grams": 137}],
        "rawMaterialSettings": {"scrapPercent": 2, "stockDate": "2026-09-21", "stocks": [{"materialCode": "MAT-T18", "kg": 901}]},
    }
    data = {"version": 1, "productionArea": "torna", "parameters": piston_parameters(), "settings": settings}
    content = build_data_package("complete", data)
    assert parse_data_package(content, "complete") == data
    workbook = load_workbook(BytesIO(content), read_only=True)
    rows = [tuple(row) for sheet in workbook if sheet.sheet_state == "visible"
            for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert any("PISTON-T18" in row and "D-01" in row and 803 in row for row in rows)
    assert any("T18 washing" in row and 480 in row and 8 in row for row in rows)
    assert any("T18 finish gap" in row and 120 in row and 2 in row for row in rows)
    assert any("MAT-T18" in row and 901 in row for row in rows)
    assert any("PISTON-T18" in row and 137 in row for row in rows)


def test_t18_complete_package_export_and_import_require_admin():
    data = {"version": 1, "productionArea": "torna", "parameters": piston_parameters(),
            "settings": {"setupSettings": {}, "machines": [], "holidays": [], "calendarEvents": []}}
    payload = {"scope": "complete", "data": data}
    admin = {"Authorization": f"Bearer {create_test_session('test-password', 'admin')}"}
    user = {"Authorization": f"Bearer {create_test_session('test-password', 'user')}"}
    with TestClient(app) as client:
        assert client.post("/api/data-packages/export", headers=user, json=payload).status_code == 403
        exported = client.post("/api/data-packages/export", headers=admin, json=payload)
        assert exported.status_code == 200
        file = {"file": ("T18-complete.xlsx", exported.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        assert client.post("/api/data-packages/complete/import", headers=user, files=file).status_code == 403
        imported = client.post("/api/data-packages/complete/import", headers=admin, files=file)
        assert imported.status_code == 200
        assert imported.json()["data"] == data
