"""Independent readable-plan checks. Saved scenarios are the sole input."""

from copy import deepcopy
from datetime import datetime, timedelta
import hashlib
from io import BytesIO
import json
from pathlib import Path

import pytest

from openpyxl import load_workbook

from app.data_package import build_data_package, parse_data_package
from app.config import settings
from app.database import connection, init_database


@pytest.fixture(autouse=True)
def isolated_scenario_database(tmp_path):
    original = settings.database_path
    object.__setattr__(settings, "database_path", tmp_path / "scenario-export.sqlite")
    try:
        init_database()
        yield
    finally:
        object.__setattr__(settings, "database_path", original)


def store_scenarios(items):
    with connection() as db:
        for item in items:
            db.execute("INSERT INTO scenarios(id,name,created_at,notes,inputs_json,result_json) VALUES(?,?,?,?,?,?)",
                       (item["id"], item["name"], item["createdAt"], item["notes"],
                        json.dumps(item["seed"]), json.dumps(item["result"])))


def scenario_fixture(identity="scenario-one", name="Aynı ad"):
    return {
        "id": identity, "name": name, "createdAt": "2026-10-18T08:00:00Z",
        "notes": "Recorded scenario; not live production",
        "seed": {
            "productionArea": "torna", "machines": [], "products": [],
            "preferences": [], "orders": [], "manualBatches": [],
            "restrictions": [], "holidays": [], "savedBatches": [],
        },
        "result": {"batches": [], "summary": {
            "demand": 0, "planned": 0, "plannedBatchCount": 0,
        }},
    }


def visible_rows(content):
    workbook = load_workbook(BytesIO(content), data_only=False)
    return [tuple(row) for sheet in workbook if sheet.sheet_state == "visible"
            for row in sheet.iter_rows(values_only=True)]


def test_t17_same_name_selections_keep_both_stable_ids_visible_and_exact_backup():
    data = [scenario_fixture("scenario-one"), scenario_fixture("scenario-two")]
    before = deepcopy(data)
    store_scenarios(data)
    content = build_data_package("scenarios", data)
    assert parse_data_package(content, "scenarios") == before
    assert data == before
    rows = visible_rows(content)
    for identity in ("scenario-one", "scenario-two"):
        assert any(identity in row and "Aynı ad" in row for row in rows), (
            "Each selected scenario needs its stable ID beside its name"
        )


def test_t17_saved_turning_work_is_readable_with_its_scenario_and_dates():
    scenario = scenario_fixture()
    batch = {
        "id": "saved-charge-t17", "workOrder": "WO-T17-203", "product": "PISTON-T17",
        "quantity": 731, "machineId": "C-17", "start": 46313.25,
        "end": 46313.75, "setupStart": 46313.2, "sequence": 1, "status": "planned",
        "orderQuantity": 731, "diameter": "17", "batchSize": 731, "dailyRate": 1462,
        "availableStart": 46313.2, "planEnd": 46343, "reason": "Saved",
        "explanation": "Saved scenario", "planColumn": 1, "existingQuantity": 0,
    }
    scenario["result"]["batches"] = [batch]
    data = [scenario]
    before = deepcopy(data)
    store_scenarios(data)
    content = build_data_package("scenarios", data)
    rows = visible_rows(content)
    matching = [row for row in rows if "WO-T17-203" in row]
    assert matching, "A summary or hidden JSON is not a readable saved Torna plan"
    assert any(all(value in row for value in (
        "scenario-one", "Aynı ad", "PISTON-T17", 731, "C-17",
    )) for row in matching)
    assert data == before
    assert parse_data_package(content, "scenarios") == before


def test_t17_empty_plan_keeps_all_four_process_sections_and_explicit_empty_status():
    data = [scenario_fixture()]
    store_scenarios(data)
    workbook = load_workbook(BytesIO(build_data_package("scenarios", data)))
    visible = [sheet for sheet in workbook if sheet.sheet_state == "visible"]
    # Section titles may vary. Test the business process labels, not sheet names.
    for process in ("Torna", "Delme", "Çapak Alma", "GKM"):
        sections = [sheet for sheet in visible if process.casefold() in (
            sheet.title + " " + " ".join(str(cell.value or "") for row in sheet for cell in row)
        ).casefold()]
        assert sections, f"An empty saved plan must still expose its {process} section"
        assert any(any(word in " ".join(str(cell.value or "") for row in sheet for cell in row).casefold()
                       for word in ("boş", "kayıt yok", "iş yok", "plan yok")) for sheet in sections), (
            f"{process} needs an explicit empty status, not fabricated production"
        )


def test_t17_untrusted_scenario_text_is_literal_excel_text_and_backup_stays_exact():
    scenario = scenario_fixture(name='=HYPERLINK("https://example.invalid","unsafe")')
    scenario["notes"] = '=1+1'
    data = [scenario]
    store_scenarios(data)
    content = build_data_package("scenarios", data)
    assert parse_data_package(content, "scenarios") == data
    workbook = load_workbook(BytesIO(content), data_only=False)
    for sheet in workbook:
        if sheet.sheet_state != "visible":
            continue
        for row in sheet:
            for cell in row:
                assert cell.data_type != "f", "User scenario text must not become an Excel formula"


def test_t17_export_does_not_invent_current_production_in_hidden_payload():
    legacy = scenario_fixture("106", "Legacy 106")
    legacy["seed"].pop("productionArea")
    legacy["seed"]["legacyUnknownField"] = {"retain": [None, False, 0, "şarj"]}
    before = deepcopy(legacy)
    content = build_data_package("scenarios", [legacy])
    assert legacy == before
    assert parse_data_package(content, "scenarios") == [before]
    # New process output may be derived for display. It must not become saved
    # current work, real WIP, or an enabled flow setting in the import payload.
    restored = parse_data_package(content, "scenarios")[0]["seed"]
    assert "flowSettings" not in restored
    assert "wipLots" not in restored
    assert "processCurrentJobs" not in restored


def full_route_scenario():
    scenario = scenario_fixture("route-scenario", "Four process plan")
    # Use only structural factory defaults. Every work item, quantity, rate and
    # resource boundary used by this test is set below.
    foundation = json.loads((Path(__file__).resolve().parents[2] / "frontend/data/planning-seed.json").read_text())
    seed = scenario["seed"]
    seed["setupSettings"] = foundation.get("setupSettings", {})
    seed["source"] = "synthetic-export-test"
    seed["schemaVersion"] = 22
    seed["products"] = [{"product": "PISTON-T17", "setupFamily": "piston", "diameter": "17",
                         "batchSize": 100, "dailyRate": 300, "shiftRate": 100}]
    machine = deepcopy(foundation["machines"][0])
    machine.update(id="C-17", active=True, availableStart=46313, planEnd=46343,
                   shiftFactor=3, capacityFactor=1, rates={"PISTON-T17": 100})
    machine["currentJob"] = {"workOrder": "", "product": "", "quantity": 0, "originalQuantity": 0,
                             "remainingQuantity": 0, "start": 46313, "end": 46313, "dailyRate": 0, "diameter": ""}
    seed["machines"] = [machine]
    seed["preferences"] = [{"key": "PISTON-T17", "machines": ["C-17"]}]
    seed["orders"] = [{"id": "order-t17", "product": "PISTON-T17", "quantity": 100,
                       "dueDate": "2026-11-30", "savedPlanned": 0, "savedDifference": 0, "savedMachines": []}]
    stages = [("drilling", "D-17"), ("deburring", "CA-17"), ("gkm", "GKM-17")]
    seed["processMasterData"] = {
        "version": 1, "source": {"kind": "synthetic", "files": [], "calculatedOnce": ""},
        "machineCodeStandard": [],
        "resources": [{"id": resource, "name": resource, "process": stage, "active": True,
                       "defaultShifts": 3, "availableStart": 46313, "planEnd": 46343, "eligibleProductGroups": ["piston"]} for stage, resource in stages],
        "products": [{"product": "PISTON-T17", "family": "piston",
                      "route": ["turning", "drilling", "deburring", "gkm"],
                      "processes": [{"process": stage, "sequence": i, "waitWorkdaysBefore": 0,
                                     "setupKey": "TEST", "resourcePriority": [resource],
                                     "unitsPerShift": {resource: 100}, "setupMinutesByResource": {resource: 0}}
                                    for i, (stage, resource) in enumerate(stages, 1)]}],
    }
    # Explicit saved Piston setup rules avoid relying on CF route fields or
    # the legacy first-job setup fallback. No transition wait is enabled.
    seed["flowSettings"] = {"version": 1, "enabled": True, "department": "piston", "defaultShifts": 3,
                            "shiftOverrides": [], "rules": [], "setupMinutes": {
                                stage: {key: 0 for key in ("first", "sameProduct", "sameKey", "differentKey", "differentFamily")}
                                for stage in ("drilling", "deburring", "gkm")}}
    seed["processWaitSettings"] = {}
    seed["wipLots"] = []
    seed["productionHistory"] = []
    seed["processCurrentJobs"] = []
    batch = {"id": "charge-t17", "workOrder": "WO-FOUR-STAGES", "product": "PISTON-T17", "quantity": 100,
             "machineId": "C-17", "start": 46313.25, "end": 46313.5, "sequence": 1, "status": "planned",
             "orderQuantity": 100, "diameter": "17", "batchSize": 100, "dailyRate": 300,
             "availableStart": 46313, "planEnd": 46343, "reason": "Synthetic", "explanation": "Saved result",
             "planColumn": 1, "existingQuantity": 0}
    scenario["result"]["batches"] = [batch]
    return scenario


def test_t17_all_four_saved_processes_conserve_quantity_and_ignore_client_date_claims():
    saved = full_route_scenario()
    store_scenarios([saved])
    client_copy = deepcopy(saved)
    client_copy["result"]["batches"][0].update(start=99999, end=100000, quantity=999)
    before = deepcopy(client_copy)
    content = build_data_package("scenarios", [client_copy])
    workbook = load_workbook(BytesIO(content))
    timeline = []
    for process, resource in [("Torna", "C-17"), ("Delme", "D-17"), ("Çapak Alma", "CA-17"), ("GKM", "GKM-17")]:
        sheet = next((sheet for sheet in workbook if process.casefold() in sheet.title.casefold()), None)
        assert sheet is not None, f"Missing readable {process} plan"
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        jobs = [row for row in rows if "WO-FOUR-STAGES" in row and resource in row]
        assert len(jobs) == 1, f"Expected one saved {process} job"
        assert all(value in jobs[0] for value in ("route-scenario", "Four process plan", "PISTON-T17", 100))
        assert 999 not in jobs[0] and 99999 not in jobs[0] and 100000 not in jobs[0]
        headers = " ".join(str(value or "") for value in rows[0]).casefold()
        for label in ("başlangıç", "bitiş", "durum", "kaynak"):
            assert label in headers, f"{process} needs an explicit {label} column"
        columns = {str(value).casefold(): index for index, value in enumerate(rows[0])}
        start, end = (jobs[0][columns[label]] for label in ("başlangıç", "bitiş"))
        assert isinstance(start, datetime) and isinstance(end, datetime) and end > start
        assert jobs[0][columns["durum"]] == "Planlandı"
        assert jobs[0][columns["kaynak türü"]] == "Torna planı"
        timeline.append((start, end))
        if process == "Torna":
            # These are saved factory wall-clock instants, not UTC plus three hours.
            assert start == datetime(1899, 12, 30) + timedelta(days=46313.25)
            assert end == datetime(1899, 12, 30) + timedelta(days=46313.5)
        else:
            assert abs((end - start).total_seconds() - 8 * 3600) < 1
    assert all(timeline[i][0] >= timeline[i - 1][1] for i in range(1, 4))
    assert client_copy == before
    assert parse_data_package(content, "scenarios") == [before]
    info = " ".join(str(cell.value or "") for row in workbook["Bilgi"] for cell in row).casefold()
    assert "sunucuda kayıtlı" in info and "motor" in info
    assert "türkiye" in info and "süreler saat" in info
    from app.scenario_plans import RUNTIME
    assert hashlib.sha256(RUNTIME.read_bytes()).hexdigest() in info
    assert "hesap zamanı" in info


def test_t17_turning_sheet_includes_stored_current_work_and_real_completed_step():
    saved = full_route_scenario()
    saved["result"]["batches"] = []
    saved["seed"]["machines"][0]["currentJob"] = {
        "workOrder": "CURRENT-T17", "product": "PISTON-T17", "quantity": 100,
        "originalQuantity": 100, "remainingQuantity": 60, "start": 46313.25,
        "end": 46313.75, "dailyRate": 300, "diameter": "17",
    }
    saved["seed"]["wipLots"] = [{
        "id": "real-lot-t17", "sourceHistoryEntryId": "real-history-t17", "sourceBatchId": "real-charge-t17",
        "workOrder": "ACTUAL-T17", "product": "PISTON-T17", "family": "piston",
        "originalQuantity": 80, "availableQuantity": 80, "deliveredQuantity": 0, "scrappedQuantity": 0,
        "stage": "drilling-queue", "nextProcess": "drilling", "readyAt": 46313.5,
        "completedSteps": [{"process": "turning", "resourceId": "C-17", "resourceName": "C-17",
                            "start": 46312.25, "end": 46312.5, "quantity": 80,
                            "recordedAt": "2026-10-17T10:00:00Z"}],
        "createdAt": "2026-10-17T10:00:00Z", "updatedAt": "2026-10-17T10:00:00Z",
    }]
    store_scenarios([saved])
    content = build_data_package("scenarios", [saved])
    workbook = load_workbook(BytesIO(content))
    sheet = next((sheet for sheet in workbook if "torna" in sheet.title.casefold()), None)
    assert sheet is not None
    rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    for work_order, quantity, status in [("CURRENT-T17", 100, "Çalışıyor"), ("ACTUAL-T17", 80, "Fiili")]:
        jobs = [row for row in rows if work_order in row]
        assert len(jobs) == 1, f"Recorded Torna work {work_order} must occur once"
        assert quantity in jobs[0] and status in jobs[0] and "C-17" in jobs[0]
    assert parse_data_package(content, "scenarios") == [saved]


@pytest.mark.parametrize("invalid", ["missing-route", "missing-resource"])
def test_t17_missing_route_or_resource_is_visible_in_warnings_not_a_claimed_complete_plan(invalid):
    saved = full_route_scenario()
    if invalid == "missing-route":
        saved["seed"]["processMasterData"]["products"] = []
    else:
        saved["seed"]["processMasterData"]["resources"] = []
    store_scenarios([saved])
    workbook = load_workbook(BytesIO(build_data_package("scenarios", [saved])))
    sheet = next((sheet for sheet in workbook if "uyarı" in sheet.title.casefold()), None)
    assert sheet is not None, "An invalid saved route needs readable warnings"
    rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    assert any("route-scenario" in row and "PISTON-T17" in row for row in rows[1:])
    for process in ("Delme", "Çapak Alma", "GKM"):
        section = next((item for item in workbook if process.casefold() in item.title.casefold()), None)
        assert section is not None and section.max_row >= 2


@pytest.fixture
def runtime_stub(tmp_path, monkeypatch):
    import hashlib
    import app.scenario_plans as plans
    runtime = tmp_path / "scenario-plan-runtime.cjs"
    runtime.write_bytes(b"// synthetic test runtime; never executed")
    runtime.with_suffix(".cjs.json").write_text(json.dumps({"runtimeSha256": hashlib.sha256(runtime.read_bytes()).hexdigest()}))
    monkeypatch.setattr(plans, "RUNTIME", runtime)
    return plans


@pytest.mark.parametrize("failure", ["timeout", "invalid-json", "oversize", "bad-schema"])
def test_t17_runtime_failure_is_explicit_bounded_and_does_not_change_saved_scenario(runtime_stub, monkeypatch, failure):
    import subprocess
    plans = runtime_stub
    saved = full_route_scenario()
    before = deepcopy(saved)
    store_scenarios([saved])
    monkeypatch.setenv("NODE_OPTIONS", "--require=/untrusted.cjs")
    monkeypatch.setattr(plans, "MAX_OUTPUT", 4096)

    def run(command, **kwargs):
        assert 0 < kwargs["timeout"] <= 30
        assert kwargs.get("shell", False) is False
        assert "NODE_OPTIONS" not in kwargs["env"]
        assert "--max-old-space-size=512" in command
        assert command[-1] == str(plans.RUNTIME)
        request = json.loads(kwargs["input"])
        assert request["scenarios"][0]["seed"] == saved["seed"]
        if failure == "timeout":
            raise subprocess.TimeoutExpired(command, kwargs["timeout"])
        if failure == "invalid-json":
            kwargs["stdout"].write(b"not-json")
        elif failure == "oversize":
            kwargs["stdout"].write(b"x" * 4097)
        else:
            kwargs["stdout"].write(b'{"version":1,"scenarios":["invalid-record"]}')
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(plans.subprocess, "run", run)
    from app.database import scenario_record
    result = plans.scenario_plans([saved], scenario_record)
    assert result["error"], "Failed runtime must not certify a computed plan"
    assert result["scenarios"][0]["status"] == "invalid"
    assert saved == before
    assert scenario_record(saved["id"])["seed"] == before["seed"]
    assert scenario_record(saved["id"])["result"] == before["result"]


def test_t17_unscheduled_saved_charge_keeps_quantity_and_has_no_invented_dates():
    saved = full_route_scenario()
    batch = saved["result"]["batches"][0]
    batch.update(status="unplanned", start=0, end=0, quantity=137, reason="Synthetic no capacity")
    store_scenarios([saved])
    workbook = load_workbook(BytesIO(build_data_package("scenarios", [saved])))
    rows = [tuple(row) for row in workbook["Torna Planı"].iter_rows(values_only=True)]
    columns = {value: index for index, value in enumerate(rows[0])}
    job = next(row for row in rows[1:] if "WO-FOUR-STAGES" in row)
    assert job[columns["Adet"]] == 137
    assert job[columns["Durum"]] == "Planlanamadı"
    assert job[columns["Başlangıç"]] is None and job[columns["Bitiş"]] is None
    assert "Synthetic no capacity" in job[columns["Not"]]


def test_t17_old_schema_has_visible_warning_without_silent_physical_data_changes():
    saved = full_route_scenario()
    saved["id"] = "106"
    saved["seed"].pop("schemaVersion")
    store_scenarios([saved])
    with connection() as db:
        before = list(db.iterdump())
    content = build_data_package("scenarios", [saved])
    workbook = load_workbook(BytesIO(content))
    warnings = [tuple(row) for row in workbook["Plan Uyarıları"].iter_rows(values_only=True)]
    assert any("106" in row and "migrated-process-data" in row for row in warnings)
    assert parse_data_package(content, "scenarios") == [saved]
    with connection() as db:
        assert list(db.iterdump()) == before


def test_t17_unsaved_client_scenario_is_labelled_and_bounded_selection_keeps_all_backup_data():
    data = [scenario_fixture(f"not-saved-{i}", f"Client scenario {i}") for i in range(51)]
    data[0]["result"]["batches"] = full_route_scenario()["result"]["batches"]
    content = build_data_package("scenarios", data)
    assert parse_data_package(content, "scenarios") == data
    workbook = load_workbook(BytesIO(content))
    rows = [tuple(row) for row in workbook["Torna Planı"].iter_rows(values_only=True)]
    first = next(row for row in rows[1:] if "not-saved-0" in row)
    assert any("Sunucuda kayıtlı değil" in str(value) for value in first)
    assert not any("WO-FOUR-STAGES" in row for row in rows)
    warnings = " ".join(str(value or "") for row in workbook["Plan Uyarıları"].iter_rows(values_only=True) for value in row)
    assert "50" in warnings and "diğerleri" in warnings


def test_t17_nonfinite_stored_input_exports_backup_with_explicit_failed_plan_warning(runtime_stub):
    saved = full_route_scenario()
    saved["seed"]["orders"][0]["quantity"] = float("nan")
    store_scenarios([saved])
    with connection() as db:
        before = list(db.iterdump())
    content = build_data_package("scenarios", [saved])
    restored = parse_data_package(content, "scenarios")
    assert json.dumps(restored, sort_keys=True) == json.dumps([saved], sort_keys=True)
    workbook = load_workbook(BytesIO(content))
    warning_text = " ".join(str(value or "") for row in workbook["Plan Uyarıları"].iter_rows(values_only=True) for value in row)
    assert "hesaplanamadı" in warning_text.casefold()
    assert not any("WO-FOUR-STAGES" in row for row in workbook["Torna Planı"].iter_rows(values_only=True))
    with connection() as db:
        assert list(db.iterdump()) == before


def test_t17_runtime_date_outside_excel_range_is_rejected_before_workbook_render(runtime_stub, monkeypatch):
    import subprocess
    plans = runtime_stub
    saved = full_route_scenario()
    store_scenarios([saved])
    row = {"id": "charge-t17", "workOrder": "WO-FOUR-STAGES", "product": "PISTON-T17", "quantity": 100,
           "resourceId": "C-17", "start": 1_000_000_000, "end": 1_000_000_001, "readyAt": None,
           "setupHours": 0, "wait": "", "source": "Torna planı", "status": "Planlandı", "note": ""}
    output = {"version": 1, "scenarios": [{"id": saved["id"], "name": saved["name"], "createdAt": saved["createdAt"],
              "status": "ok", "rows": {"turning": [row], "drilling": [], "deburring": [], "gkm": []}, "warnings": []}]}

    def run(command, **kwargs):
        kwargs["stdout"].write(json.dumps(output).encode())
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(plans.subprocess, "run", run)
    content = build_data_package("scenarios", [saved])
    assert parse_data_package(content, "scenarios") == [saved]
    workbook = load_workbook(BytesIO(content))
    warning_text = " ".join(str(value or "") for row in workbook["Plan Uyarıları"].iter_rows(values_only=True) for value in row)
    assert "hesaplanamadı" in warning_text.casefold()
    assert not any("WO-FOUR-STAGES" in values for values in workbook["Torna Planı"].iter_rows(values_only=True))


def test_t17_deferred_cf_export_keeps_its_existing_summary_and_backup_shape(monkeypatch):
    from app.production_area import production_area
    import app.scenario_plans as plans
    scenario = scenario_fixture("cf-scenario", "CF saved scenario")
    scenario["seed"]["productionArea"] = "cubuk-filtre"
    calls = []

    def forbidden_runtime(*args, **kwargs):
        calls.append(True)
        raise AssertionError("Piston readable export must not run on deferred CF")

    monkeypatch.setattr(plans, "_run", forbidden_runtime)
    token = production_area.set("cubuk-filtre")
    try:
        content = build_data_package("scenarios", [scenario])
    finally:
        production_area.reset(token)
    workbook = load_workbook(BytesIO(content))
    assert set(workbook.sheetnames) == {"Bilgi", "Senaryolar", "__SelsaVeri"}
    assert parse_data_package(content, "scenarios") == [scenario]
    assert not calls
