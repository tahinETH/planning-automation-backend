from copy import deepcopy
from io import BytesIO
import pytest
from openpyxl import load_workbook
from app.raw_materials import initial_raw_material_settings, validate_raw_material_settings
from app.database import init_database, connection, planning_state, save_planning_state, ProtectedSettingsChange
from app.data_package import build_data_package, parse_data_package
from app.ravi.tools import Inspect, inspect_state
from app.overview_export import build_overview_workbook
from app.models import OverviewOperationPlan


def test_material_settings_migration_permissions_validation_and_legacy_saves():
    init_database()
    with connection() as db:
        db.execute("DELETE FROM planning_state")
    save_planning_state({"orders": []}, can_manage_settings=True)
    init_database()
    state = planning_state()
    assert state["seed"]["rawMaterialSettings"] == initial_raw_material_settings()
    settings = {"scrapPercent": 5, "stockDate": "2026-09-14", "stocks": [{"materialCode": "Z1", "kg": 51}]}
    incoming = deepcopy(state["seed"])
    incoming["rawMaterialSettings"] = settings
    with pytest.raises(ProtectedSettingsChange):
        save_planning_state(incoming, state["updatedAt"], can_manage_settings=False)
    state = save_planning_state(incoming, state["updatedAt"], can_manage_settings=True)
    init_database()
    assert planning_state()["seed"]["rawMaterialSettings"] == settings
    legacy = deepcopy(state["seed"])
    del legacy["rawMaterialSettings"]
    assert save_planning_state(legacy, state["updatedAt"], can_manage_settings=False)["seed"]["rawMaterialSettings"] == settings
    for rate in [-1, 101, True, float("nan"), float("inf"), "2", None]:
        with pytest.raises(ValueError):
            validate_raw_material_settings({**settings, "scrapPercent": rate})
    for stock_date in ["2026-02-30", "20260914", "bad", None]:
        with pytest.raises(ValueError):
            validate_raw_material_settings({**settings, "stockDate": stock_date})
    for kg in [-1, True, float("inf"), "51"]:
        with pytest.raises(ValueError):
            validate_raw_material_settings({**settings, "stocks": [{"materialCode": "Z1", "kg": kg}]})
    validate_raw_material_settings({**settings, "scrapPercent": 0, "stocks": [{"materialCode": "Z1", "kg": None}]})
    with pytest.raises(ValueError):
        validate_raw_material_settings({**settings, "stocks": settings["stocks"] * 2})


def test_material_backup_and_ravi_show_units_and_snapshot_without_mutation():
    data = {"rawMaterialSettings": {"scrapPercent": 5, "stockDate": "2026-09-14", "stocks": [{"materialCode": "Z1", "kg": 0}, {"materialCode": "Z2", "kg": None}]}}
    content = build_data_package("settings", data)
    assert parse_data_package(content, "settings")["rawMaterialSettings"] == data["rawMaterialSettings"]
    sheet = load_workbook(BytesIO(content))["Hammadde Stokları"]
    assert sheet["B2"].value == 0
    snapshot = {"seed": data, "updatedAt": "2026-09-14"}
    before = deepcopy(snapshot)
    result = inspect_state(snapshot, Inspect(collection="raw_material_stocks", query="Z1"))
    assert result["total"] == 1
    assert result["rows"][0]["Ambar stoğu (kg)"] == 0
    assert result["rows"][0]["Iskarta Oranı (%)"] == 5
    assert "kg" in result["note"]
    assert snapshot == before


@pytest.mark.parametrize("process", ["turning", "drilling"])
def test_adjustable_scrap_is_used_in_export_cells_and_labels(process):
    plan = {"process": process, "label": "Torna" if process == "turning" else "Delme", "scrapPercent": 5,
            "totalQuantity": 1000, "totalJobCount": 1, "resources": [{"id": "C-01", "name": "Test", "rows": [
                {"status": "planned", "position": 1, "product": "P1", "quantity": 1000, "materialCode": "Z1", "unitWeightGrams": 50}]}]}
    validated = OverviewOperationPlan.model_validate(plan).model_dump()
    content = build_overview_workbook({"selectedProcess": process, "operationPlans": [validated], "generatedAt": "2026-09-14", "dirty": False, "printRange": {"mode": "next-jobs"}})
    workbook = load_workbook(BytesIO(content))
    values = [cell.value for sheet in workbook for row in sheet for cell in row]
    assert 52.5 in values
    assert any("%5 ıskarta" in str(value) for value in values)
    assert not any("%2 ıskarta" in str(value) for value in values)
